import logging

import openpyxl
from pathlib import Path


logger = logging.getLogger(__name__)

def analyze_excel(file_path: str):
    """Analyze the Excel file structure using openpyxl."""
    wb = openpyxl.load_workbook(file_path)
    sheet = wb.active

    logger.info(f"Excel file: {file_path}")
    logger.info(f"Number of rows: {sheet.max_row}")
    logger.info(f"Number of columns: {sheet.max_column}")

    logger.info("\nFirst 10 rows:")
    for row_idx in range(1, min(11, sheet.max_row + 1)):
        row_values = []
        for col_idx in range(1, sheet.max_column + 1):
            cell_value = sheet.cell(row=row_idx, column=col_idx).value
            row_values.append(cell_value)
        logger.info(f"  Row {row_idx}: {' | '.join([str(v) or '' for v in row_values])}")

    wb.close()


if __name__ == "__main__":
    import sys

    if len(sys.argv) < 2:
        logger.info("Usage: python excel_analysis.py <excel_file.xlsx>")
    else:
        analyze_excel(sys.argv[1])
